from pathlib import Path
from src.report import process_files, export_report
from openpyxl import load_workbook


def test_sample_pipeline(tmp_path):
    folder = Path(__file__).resolve().parents[1] / 'data'
    clean, rejected, duplicates = process_files(folder)
    assert len(clean) == 8
    assert len(rejected) == 1
    assert duplicates == 1
    out = tmp_path / 'report.xlsx'
    export_report(clean, rejected, duplicates, out)
    book = load_workbook(out)
    assert {'Executive Summary', 'Cleaned Orders', 'Rejected Rows'} <= set(book.sheetnames)
    assert book['Executive Summary']['B3'].value == 8
    assert book['Executive Summary']['B4'].value == 48250
