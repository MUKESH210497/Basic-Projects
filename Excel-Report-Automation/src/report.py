"""Consolidate monthly sales CSV files into a formatted Excel report."""
from pathlib import Path
import argparse
import pandas as pd
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.chart import BarChart, Reference

REQUIRED = {'order_id', 'date', 'region', 'product', 'quantity', 'unit_price'}


def process_files(folder: Path):
    files = sorted(folder.glob('*.csv'))
    if not files:
        raise ValueError('No CSV files found')
    tables = []
    for file in files:
        df = pd.read_csv(file, dtype={'order_id': str})
        missing = REQUIRED - set(df.columns)
        if missing:
            raise ValueError(f'{file.name}: missing {sorted(missing)}')
        tables.append(df[list(sorted(REQUIRED))])
    raw = pd.concat(tables, ignore_index=True)
    before = len(raw)
    raw = raw.drop_duplicates(subset=['order_id'], keep='first').copy()
    raw['date'] = pd.to_datetime(raw['date'], errors='coerce')
    raw['quantity'] = pd.to_numeric(raw['quantity'], errors='coerce')
    raw['unit_price'] = pd.to_numeric(raw['unit_price'], errors='coerce')
    valid = raw['date'].notna() & raw['region'].notna() & raw['product'].notna() & raw['order_id'].notna() & (raw['quantity'] > 0) & (raw['unit_price'] >= 0)
    rejected = raw.loc[~valid].copy()
    clean = raw.loc[valid].copy()
    clean['revenue'] = (clean['quantity'] * clean['unit_price']).round(2)
    clean = clean.sort_values(['date', 'order_id'])
    return clean, rejected, before - len(raw)


def export_report(clean, rejected, duplicate_count, output: Path):
    wb = Workbook()
    ws = wb.active
    ws.title = 'Executive Summary'
    ws.append(['Sales Report | Synthetic Sample Data'])
    ws.append(['Metric', 'Value'])
    ws.append(['Valid orders', len(clean)])
    ws.append(['Total revenue', round(float(clean['revenue'].sum()), 2)])
    ws.append(['Rejected rows', len(rejected)])
    ws.append(['Duplicate IDs removed', duplicate_count])
    ws.append([])
    ws.append(['Region', 'Revenue'])
    regions = clean.groupby('region')['revenue'].sum().sort_values(ascending=False)
    for region, revenue in regions.items():
        ws.append([region, round(float(revenue), 2)])
    if len(regions):
        chart = BarChart()
        chart.title = 'Revenue by Region'
        chart.y_axis.title = 'Revenue (sample currency units)'
        chart.add_data(Reference(ws, min_col=2, min_row=8, max_row=8 + len(regions)), titles_from_data=True)
        chart.set_categories(Reference(ws, min_col=1, min_row=9, max_row=8 + len(regions)))
        ws.add_chart(chart, 'D8')
    for sheet_name, df in [('Cleaned Orders', clean), ('Rejected Rows', rejected)]:
        sh = wb.create_sheet(sheet_name)
        sh.append(list(df.columns))
        for row in df.itertuples(index=False, name=None):
            sh.append([v.to_pydatetime() if isinstance(v, pd.Timestamp) else v for v in row])
        sh.freeze_panes = 'A2'
        sh.auto_filter.ref = sh.dimensions
        for cell in sh[1]:
            cell.fill = PatternFill('solid', fgColor='17365D')
            cell.font = Font(color='FFFFFF', bold=True)
        for col in sh.columns:
            sh.column_dimensions[col[0].column_letter].width = min(30, max(14, max(len(str(c.value or '')) for c in col) + 2))
    ws['A1'].font = Font(size=16, bold=True, color='17365D')
    for rownum in (2, 8):
        for cell in ws[rownum]:
            cell.fill = PatternFill('solid', fgColor='17365D')
            cell.font = Font(color='FFFFFF', bold=True)
    ws.column_dimensions['A'].width = 30
    ws.column_dimensions['B'].width = 22
    ws['B4'].number_format = '#,##0.00'
    for row in range(9, 9 + len(regions)):
        ws.cell(row, 2).number_format = '#,##0.00'
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output)
    return output


def main():
    parser = argparse.ArgumentParser()
    parser.add_argument('--input', type=Path, default=Path('data'))
    parser.add_argument('--output', type=Path, default=Path('outputs/sales_report.xlsx'))
    args = parser.parse_args()
    clean, rejected, duplicates = process_files(args.input)
    export_report(clean, rejected, duplicates, args.output)
    print(f'Created {args.output} | valid={len(clean)} rejected={len(rejected)} duplicates={duplicates}')


if __name__ == '__main__':
    main()
